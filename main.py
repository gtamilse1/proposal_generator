import os
import logging
import re
import argparse
from docx import Document
from datetime import datetime

"""
Project Overview:
This project aims to automate the process of generating customized MS Word documents. The main steps involved are:

1. Ingest MS Word documents: Read and parse the content of existing MS Word documents.
2. Identify all variables: Detect placeholders or variables within the document that need to be populated with data.
3. Populate variables with data: Replace the identified placeholders with actual data from a given source.
4. Generate a new MS Word document: Create a new document with the populated data, ready for use.

This automation will streamline the process of creating personalized documents, saving time and reducing errors.
"""
# Set up logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("app.log"),
        logging.StreamHandler()
    ]
)

def usr_input_val():
    """
    This function handles user input validation for the architecture type.
    It first attempts to parse the architecture from command-line arguments.
    If not provided, it prompts the user to input the architecture.
    The function then validates the input to ensure it is a string.
    Returns the validated architecture string or None if validation fails.
    """
    # Check if architecture is defined
    architecture = None
    
    # Parse command-line arguments
    parser = argparse.ArgumentParser(description='Process architecture input.')
    parser.add_argument('-a', '--architecture', type=str, help='Architecture type')
    args = parser.parse_args()
    
    if args.architecture:
        architecture = args.architecture.strip().lower()
    
    # Prompt for input if architecture is not defined
    if not architecture:
        architecture = input("Please enter the architecture: ").strip().lower()
    
    # Validate user input
    if not isinstance(architecture, str):
        logging.error("Invalid input. Please enter a valid string for architecture.")
        # print("Invalid input. Please enter a valid string for architecture.")
        return None

    return architecture

def classify_documents(directory, architecture):
    """
    Classifies documents in a given directory based on the specified architecture.
    This function searches for a .docx file in the specified directory that contains
    the architecture keyword in its filename. If such a file is found, it attempts to
    load the document and extract variables enclosed in double angle brackets (e.g., <<variable>>).
    The function returns the filename and a list of unique variables found in the document.
    Args:
        directory (str): The path to the directory containing the documents.
        architecture (str): The architecture keyword to search for in the filenames.
    Returns:
        tuple: A tuple containing the filename (str) and a list of unique variables (list of str).
               If no document is found or an error occurs, returns (None, None).
    Raises:
        None: This function handles exceptions internally and logs errors.
    """

    matching_files = []
    for filename in os.listdir(directory):
        if filename.endswith(".docx") and architecture in filename.lower():
            file_to_load = os.path.join(directory, filename)
            matching_files.append(file_to_load)

    if not matching_files:
        logging.error(f"No document found for architecture: {architecture}")
        print(f"No document found for architecture: {architecture}")
        return None, None

    # Select the latest file based on modification time
    latest_file = max(matching_files, key=os.path.getmtime)

    try:
        doc = Document(latest_file)
        logging.info(f"Loaded document: {latest_file}")
    except Exception as e:
        logging.error(f"Failed to read document: {latest_file}, error: {e}")
        print(f"Failed to read document: {latest_file}, error: {e}")
        return None, None

    doc_text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
    variables = re.findall(r'<<(\w+)>>', doc_text)
    file_vars = list(set(variables))  # De-duplicate the list of variables

    if file_vars:
        logging.info(f"Found variables in {architecture} document: {file_vars}")
    else:
        logging.info(f"No variables found in {architecture} document")

    return latest_file, file_vars

def main():
    directory = '/home/cisco/bdm-genai/GIT/gtamilse/proposal_generator/ptemplates'
    
    # Get validated user input for architecture
    architecture = usr_input_val()
    if not architecture:
        return

    # Classify documents and find variables
    latest_file, file_vars = classify_documents(directory, architecture)

    if latest_file and file_vars is not None:
        print(f"Filename: {latest_file}")
        print(f"Variables: {file_vars}")
    else:
        print("No document or variables found.")

if __name__ == "__main__":
    main()